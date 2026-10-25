"""Motor 1 (S3): OpenPyXL baseline, Python puro, modo read_only.

Sin dependencias compiladas: referencia de consumo plano y latencia peor caso.
En read_only el parseo XML ocurre durante la iteración, así que I/O aquí =
apertura + streaming de filas; CPU = agregación del dict en memoria.
"""

from __future__ import annotations

from ._child import run_child


def load(path: str) -> list[tuple]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["sales"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    wb.close()
    return rows


def aggregate(rows: list[tuple]) -> dict:
    totals: dict[str, float] = {}
    for _, suc, monto, _ in rows:
        totals[suc] = totals.get(suc, 0.0) + float(monto or 0.0)
    return totals


if __name__ == "__main__":
    run_child("openpyxl", load, aggregate)
